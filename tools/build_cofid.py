"""Builds data/cofid.json from the official UK food composition dataset (CoFID, McCance and
Widdowson, Open Government Licence). Run by .github/workflows/build-cofid.yml, because the
gov.uk download is easiest from a GitHub runner. Output is one compact row per food:
  { c: food code, n: name, g: group code, kcal, prot, fat, sat, carb, sugar, starch, fibre, salt }
All per 100 g. Missing values are null. "Tr" (trace) counts as 0. """
import io, json, os, re, sys, urllib.request

PAGE = 'https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid'
OUT = 'data/cofid.json'
UA = {'User-Agent': 'Mozilla/5.0 (care-log build script)'}

def get(url):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()

html = get(PAGE).decode('utf-8', 'replace')
links = sorted(set(re.findall(r'href="(https://assets\.publishing\.service\.gov\.uk/[^"]+\.xlsx)"', html)))
print('xlsx links found:', *links, sep='\n  ')
if not links:
    sys.exit('No xlsx links found on the publication page')

best, best_bytes = None, b''
for u in links:
    b = get(u)
    print(f'{u} -> {len(b)} bytes')
    if len(b) > len(best_bytes):
        best, best_bytes = u, b
print('Using', best)

import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(best_bytes), read_only=True, data_only=True)
print('sheets:', wb.sheetnames)

def find_sheet(word):
    for name in wb.sheetnames:
        if word.lower() in name.lower():
            return wb[name]
    return None

def header_map(ws):
    """Returns (header row index, {normalised header: column index}) using the first row that has 'Food Code'."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
        cells = [str(c).strip() if c is not None else '' for c in row]
        if any('food code' in c.lower() for c in cells):
            return i + 1, {re.sub(r'\s+', ' ', c.lower()): j for j, c in enumerate(cells) if c}
    raise SystemExit(f'No header row with "Food Code" in sheet {ws.title}')

def col(hmap, *needles):
    for key, j in hmap.items():
        if all(n in key for n in needles):
            return j
    return None

def num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return round(float(v), 2)
    s = str(v).strip()
    if s in ('', 'N', 'N/A'): return None
    if s.lower() in ('tr', 'trace'): return 0.0
    try: return round(float(s.replace(',', '')), 2)
    except ValueError: return None

prox = find_sheet('proximates')
inorg = find_sheet('inorganics')
if prox is None: sys.exit('No Proximates sheet')
hrow, hm = header_map(prox)
print('proximates headers:', list(hm.keys()))
ci = {
    'code': col(hm, 'food code'), 'name': col(hm, 'food name'), 'group': col(hm, 'group'),
    'kcal': col(hm, 'energy', 'kcal'), 'prot': col(hm, 'protein'), 'fat': col(hm, 'fat (g)'),
    'carb': col(hm, 'carbohydrate'), 'sugar': col(hm, 'total sugars'), 'starch': col(hm, 'starch'),
    'fibre': col(hm, 'aoac'), 'sat': col(hm, 'satd', '100g fd') or col(hm, 'saturated', 'fd'),
}
print('column picks:', ci)
missing = [k for k, v in ci.items() if v is None and k not in ('sat', 'starch')]
if missing: sys.exit('Missing columns: ' + ', '.join(missing))

foods = {}
order = []
for row in prox.iter_rows(min_row=hrow + 1, values_only=True):
    code = row[ci['code']]
    if not code or not str(code).strip(): continue
    code = str(code).strip()
    name = row[ci['name']]
    if not name: continue
    rec = {'c': code, 'n': str(name).strip(), 'g': (str(row[ci['group']]).strip() if row[ci['group']] else None)}
    for k in ('kcal', 'prot', 'fat', 'sat', 'carb', 'sugar', 'starch', 'fibre'):
        j = ci[k]
        rec[k] = num(row[j]) if j is not None else None
    rec['salt'] = None
    foods[code] = rec
    order.append(code)

try:
    if inorg is None: raise SystemExit('no inorganics sheet')
    hrow2, hm2 = header_map(inorg)
    jcode, jna = col(hm2, 'food code'), col(hm2, 'sodium')
    print('inorganics headers:', list(hm2.keys())[:12], '... sodium column:', jna)
    if jna is None: raise SystemExit('no sodium column')
except SystemExit as e:
    print('WARNING: salt not available:', e)
    jna = None
if jna is not None:
    if True:
        for row in inorg.iter_rows(min_row=hrow2 + 1, values_only=True):
            code = row[jcode]
            if code and str(code).strip() in foods:
                na = num(row[jna])
                foods[str(code).strip()]['salt'] = round(na * 2.5 / 1000, 2) if na is not None else None

rows = [foods[c] for c in order]
os.makedirs('data', exist_ok=True)
with open(OUT, 'w', encoding='utf-8') as f:
    json.dump({'source': 'CoFID 2021, McCance and Widdowson, Public Health England, Open Government Licence v3', 'url': best, 'count': len(rows), 'per': '100 g', 'foods': rows}, f, ensure_ascii=False, separators=(',', ':'))
print('wrote', OUT, len(rows), 'foods,', os.path.getsize(OUT), 'bytes')
for r in rows:
    if re.search(r'^(bananas, flesh|almonds, flaked|broccoli, green, boiled|chicken, breast, grilled)', r['n'], re.I):
        print(r)
